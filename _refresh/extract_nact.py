"""SNID activity countdown calendar -> const NACT (planned timeline vs NEAP status).

The countdown-of-SNIDs workbook is a manual Gantt grid, one sheet per campaign:
row 1 carries the month group headers (October / November), row 2 the day numbers
(cols C..AC = Oct 5..31, AD..AS = Nov 1..16 in the Nov-2026 sheet) and rows 4..31
the numbered activities. Coloured cell fills mark each activity's window:

  accent4 (theme 7)      main activity bars
  FF00B0F0               secondary bars (correction / re-do activities)
  FFFF0000               post-campaign review bars
  accent2 (theme 5)      catch-up days and LQAS/PCM (text says which)
  accent6 (theme 9)      campaign week start cell
  theme 0 (white) / FFFFC000 (Sundays)  grid base - never bars

Text blocks merged inside the grid (DPEC submission note, MPQA, training FLAGS,
campaign schedule) travel as `x` activity notes. The campaign date, title and
vaccine are read from the B1 header; the campaign month is asserted against the
active scope (sources.json -> auto.snd.config.start) so a stale calendar cannot
be embedded after the scope moves to the next round.

Usage: python _refresh/extract_nact.py [--out PATH]
Prints: OK nact <json chars> <summary> | ERR nact <reason>
"""
import argparse
import datetime
import os
import re
import sys

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import srcutil

DS = 'nact'
PKT = datetime.timezone(datetime.timedelta(hours=5), 'PKT')
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']

TITLES = [
    'Previous SIA action points review',
    'PEOC / divisional staff orientation',
    'District master-trainers training (PEOC)',
    'Post-campaign review + DPEC meeting',
    'Zero-dose list collection + micro plan update',
    'AIC & UCMO one-day training (micro plan updating)',
    'MP desk review + field validation (UC staff)',
    'UPEC meeting + summary micro plan to DEOC',
    'IDIMS data entry + target review',
    'Communication & social mobilization',
    'Real-time MPQA of selected UCs',
    'Course correction for failed UCs',
    'Re-MPQA if required',
    'FLW training (stall methodology)',
    'Training of defaulters',
    'Vaccine & logistics distribution',
    'District readiness meeting (DC chairs)',
    'Fixing pre-campaign issues',
    'Campaign implementation (bOPV)',
    'Market survey',
    'Zero-dose list to vaccinator',
    'Post-campaign evaluation (LQAS + PCM)',
    'LQAS/PCM investigation + UC improvement plan',
    'Retrieval of markers / vaccine / carriers / ID cards',
    'Post-campaign review at UC level',
    'Post-campaign review at district level',
    'Post-campaign review at provincial level',
    'Post-campaign review at national level',
]


def is_bar_fill(f):
    """kind letter for an activity-bar fill, or None for the grid base."""
    if f is None or f.patternType is None:
        return None
    sc = f.start_color
    if getattr(sc, 'type', None) == 'rgb':
        rgb = getattr(sc, 'rgb', None)
        if rgb == 'FFFFC000':
            return None            # Sunday
        if rgb == 'FF00B0F0':
            return 's'
        if rgb == 'FFFF0000':
            return 'r'
        return None
    if getattr(sc, 'type', None) == 'theme':
        th = getattr(sc, 'theme', None)
        return {7: 'm', 5: 'u', 9: 'c'}.get(th)   # theme 0 = white base
    return None


def read_calendar(path):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = str(ws.cell(row=1, column=2).value or '')
    m = re.search(r'\(([^)]+)\)\s*(\d{1,2})\s*-\s*(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})', hdr)
    if not m:
        raise RuntimeError('cannot read campaign dates from B1 header %r' % hdr[:80])
    vac, d0, d1, mon, year = m.group(1), int(m.group(2)), int(m.group(3)), m.group(4), int(m.group(5))
    mo = None
    for i, x in enumerate(MONTHS):
        if mon.lower().startswith(x.lower()):
            mo = i + 1
    if mo is None:
        raise RuntimeError('unknown month %r in B1 header' % mon)
    st = '%04d-%02d-%02d' % (year, mo, d0)
    en = '%04d-%02d-%02d' % (year, mo, d1)
    cm = re.search(r'([A-Za-z]+)\s+(\d{4})', hdr)
    camp = 'SNID %s-%02d' % (MONTHS[mo - 1], year % 100) if cm else 'SNID'

    # column -> date: row 2 day numbers, cols 3..29 = the month before the
    # campaign month, cols 30..45 = the campaign month.
    dates, sundays = {}, set()
    for col in range(3, 46):
        v = ws.cell(row=2, column=col).value
        if v is None:
            raise RuntimeError('no day number in row 2 col %d' % col)
        if col <= 29:
            mm = mo - 1 if mo > 1 else 12
            yy = year if mo > 1 else year - 1
        else:
            mm, yy = mo, year
        dates[col] = '%04d-%02d-%02d' % (yy, mm, int(v))

    acts = []
    for r in range(4, 32):
        n = ws.cell(row=r, column=1).value
        name = ws.cell(row=r, column=2).value
        if n is None or name is None:
            continue
        if int(n) != len(acts) + 1:
            raise RuntimeError('activity numbering broken at row %d (%r)' % (r, n))
        bars, kinds = [], []
        for col in range(3, 46):
            c = ws.cell(row=r, column=col)
            k = is_bar_fill(c.fill)
            if not k:
                if c.fill is None or c.fill.patternType is None:
                    continue
                sc = c.fill.start_color
                if getattr(sc, 'type', None) == 'rgb' and getattr(sc, 'rgb', None) == 'FFFFC000':
                    sundays.add(dates[col])
                continue
            if k == 'u' and n == 22:
                k = 'q'            # LQAS/PCM day on the post-evaluation row
            if kinds and col == kinds[-1][0] + 1 and k == kinds[-1][1]:
                bars[-1][1] = dates[col]
            else:
                bars.append([dates[col], dates[col], k])
            kinds.append((col, k))
        if not bars:
            raise RuntimeError('activity %d (row %d) has no bar cells' % (n, r))
        for b in bars:
            if b[2] == 'c':
                b[1] = en     # the campaign cell is drawn on day 1; the round runs to `en`
        acts.append({'n': int(n), 't': TITLES[int(n) - 1],
                     'd': re.sub(r'\s+', ' ', str(name)).strip(),
                     's': bars[0][0], 'e': bars[-1][1], 'b': bars})

    # merged text blocks inside the grid become activity notes
    for mr in ws.merged_cells.ranges:
        if mr.min_row < 4 or mr.min_row > 31 or mr.min_col < 3:
            continue
        v = ws.cell(row=mr.min_row, column=mr.min_col).value
        if v is None or not str(v).strip():
            continue
        acts[mr.min_row - 4]['x'] = re.sub(r'\s+', ' ', str(v)).strip()
    # plain (unmerged) text cells too ('Catch Up Day', '3+1 (MT)...')
    for r in range(4, 32):
        for col in range(3, 46):
            c = ws.cell(row=r, column=col)
            if c.value is None or not str(c.value).strip():
                continue
            if any(c.coordinate in mr for mr in ws.merged_cells.ranges):
                continue
            a = acts[r - 4]
            txt = re.sub(r'\s+', ' ', str(c.value)).strip()
            a['x'] = (a.get('x') + ' ' + txt).strip() if a.get('x') else txt

    # sanity: the orange fill must be exactly the Sundays spanning the grid
    d = datetime.date.fromisoformat(min(dates.values()))
    d1 = datetime.date.fromisoformat(max(dates.values()))
    d += datetime.timedelta(days=(6 - d.weekday()) % 7)   # weekday(): Mon=0..Sun=6
    exp_sun = []
    while d <= d1:
        exp_sun.append(d.isoformat())
        d += datetime.timedelta(days=7)
    if sundays != set(exp_sun):
        raise RuntimeError('Sunday (orange) columns %s do not match %s - day numbering shifted?'
                           % (sorted(sundays), sorted(exp_sun)))
    if len(acts) != 28:
        raise RuntimeError('expected 28 activities, found %d' % len(acts))
    return camp, vac, st, en, acts


def build(cfg):
    e = cfg['auto'][DS]
    files = srcutil.discover(e)
    if not files:
        raise SystemExit('ERR %s no activity calendar under %s' % (DS, e['roots']))
    rec = files[0]
    camp, vac, st, en, acts = read_calendar(rec['path'])
    snd = (cfg['auto'].get('snd') or {}).get('config') or {}
    if snd.get('start') and snd['start'][:7] != st[:7]:
        raise SystemExit('ERR %s calendar is %s (%s) but the active scope campaign starts %s - '
                         'update auto.snd.config or drop the right calendar' % (DS, camp, st, snd['start']))
    asof = datetime.datetime.fromtimestamp(rec['mtime'], PKT).strftime('%Y-%m-%d')
    out = {'camp': camp, 'vac': vac, 'st': st, 'en': en,
           'from': min(a['s'] for a in acts), 'to': max(a['e'] for a in acts),
           'asof': asof, 'src': rec['name'], 'a': acts}
    nb = sum(len(a['b']) for a in acts)
    return out, '%d activities, %d bars, %s..%s, campaign %s%s' % (
        len(acts), nb, out['from'], out['to'], st, ' (' + vac + ')' if vac else '')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--source')
    ap.add_argument('--out')
    a = ap.parse_args()
    cfg = srcutil.load_sources()
    src = a.source
    if src and not (cfg['auto'][DS].get('roots')):
        raise SystemExit('ERR %s no roots configured' % DS)
    if src:
        cfg['auto'][DS] = dict(cfg['auto'][DS], roots=[os.path.dirname(os.path.abspath(src))],
                               patterns=[os.path.basename(src)], exclude=[])
    out, summary = build(cfg)
    js = srcutil.dumps(out)
    dest = a.out or os.path.join(srcutil.OUT, DS + '.json')
    srcutil.write_text(dest, js)
    print('OK %s %d %s' % (DS, len(js), summary))


if __name__ == '__main__':
    main()
